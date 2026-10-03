#!/usr/bin/env python3
"""Gera as versões DOCX e PDF da resenha crítica.

Uso:
    python -m pip install -r requirements.txt
    python scripts/generate_deliverables.py
"""

from __future__ import annotations

import html
import re
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "RESENHA_CRITICA.md"
OUTPUT_DIR = ROOT / "entrega"
COVER_PNG = ROOT / "assets" / "capa-documento.png"
DOCX_PATH = OUTPUT_DIR / "Resenha_Critica_FBI_Sentinel.docx"
PDF_PATH = OUTPUT_DIR / "Resenha_Critica_FBI_Sentinel.pdf"

NAVY = "17233F"
BLUE = "3559A8"
ORANGE = "E85D3F"
TEAL = "2C8C7C"
SLATE = "66738D"
LIGHT = "F2F5FA"
WHITE = "FFFFFF"


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    number: str = ""


INLINE_RE = re.compile(
    r"\*\*(.+?)\*\*|\*([^*]+?)\*|\[([^\]]+)\]\((https?://[^)]+)\)|"
    r"<(https?://[^>]+)>|\[\^(\d+)\]"
)


def parse_markdown(path: Path) -> list[Block]:
    """Converte o subconjunto de Markdown usado na resenha em blocos."""
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    paragraph: list[str] = []
    started = False

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append(Block("paragraph", " ".join(x.strip() for x in paragraph)))
            paragraph = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("## Resumo"):
            started = True
        if not started:
            continue

        if not line.strip():
            flush()
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), len(heading.group(1)) - 1))
            continue

        if line.strip() == "---":
            flush()
            blocks.append(Block("rule"))
            continue

        reference = re.match(r"^\[\^(\d+)\]:\s*(.+)$", line)
        if reference:
            flush()
            blocks.append(Block("reference", reference.group(2), number=reference.group(1)))
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        if bullet:
            flush()
            blocks.append(Block("bullet", bullet.group(1)))
            continue

        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if numbered:
            flush()
            blocks.append(Block("number", numbered.group(2), number=numbered.group(1)))
            continue

        paragraph.append(line)

    flush()
    return blocks


def find_font(*names: str) -> Path:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/usr/share/fonts/truetype/liberation2"),
        Path("/usr/share/fonts/truetype/liberation"),
    ]
    for folder in candidates:
        for name in names:
            path = folder / name
            if path.exists():
                return path
    raise FileNotFoundError(f"Fonte não encontrada: {', '.join(names)}")


def generate_cover_png(path: Path) -> None:
    """Cria uma capa A4 autoral usada na versão Word."""
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1240, 1754
    image = Image.new("RGB", (width, height), "#10182e")
    draw = ImageDraw.Draw(image)

    # Gradiente vertical discreto.
    top = (16, 24, 46)
    bottom = (32, 51, 90)
    for y in range(height):
        ratio = y / (height - 1)
        color = tuple(round(top[i] * (1 - ratio) + bottom[i] * ratio) for i in range(3))
        draw.line((0, y, width, y), fill=color)

    # Geometria decorativa original.
    draw.ellipse((825, -150, 1425, 450), fill="#2a416c")
    draw.ellipse((900, 1130, 1530, 1760), fill="#101a34")
    draw.rounded_rectangle((90, 110, 395, 174), radius=32, fill="#2c8c7c")

    regular = find_font("DejaVuSans.ttf", "LiberationSans-Regular.ttf")
    bold = find_font("DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf")
    font_tag = ImageFont.truetype(str(bold), 26)
    font_kicker = ImageFont.truetype(str(bold), 28)
    font_title = ImageFont.truetype(str(bold), 83)
    font_subtitle = ImageFont.truetype(str(regular), 50)
    font_description = ImageFont.truetype(str(regular), 29)
    font_meta = ImageFont.truetype(str(regular), 23)
    font_badge = ImageFont.truetype(str(bold), 57)

    draw.text((242, 142), "ESTUDO DE CASO", font=font_tag, fill="white", anchor="mm")
    draw.text((90, 295), "SCRUM  •  AGILE  •  GOVERNO", font=font_kicker, fill="#ff8a5b")
    draw.text((90, 405), "Scrum no FBI", font=font_title, fill="white")
    draw.text((90, 515), "Projeto Sentinel", font=font_subtitle, fill="#d6e0f3")
    draw.rounded_rectangle((90, 610, 740, 618), radius=4, fill="#e85d3f")
    draw.text((90, 680), "Uma análise crítica sobre agilidade,", font=font_description, fill="#b9c7df")
    draw.text((90, 728), "governança e transformação digital", font=font_description, fill="#b9c7df")

    # Símbolo abstrato: incremento + ciclo de feedback.
    cx, cy = 920, 910
    draw.ellipse((cx - 180, cy - 180, cx + 180, cy + 180), fill="#f7f8fb")
    draw.ellipse((cx - 145, cy - 145, cx + 145, cy + 145), outline="#dce3ef", width=4)
    draw.rounded_rectangle((cx - 72, cy - 100, cx + 72, cy + 110), radius=58, outline="#17233f", width=18)
    draw.line((cx - 35, cy - 42, cx + 40, cy - 42), fill="#e85d3f", width=16)
    draw.line((cx - 35, cy + 5, cx + 40, cy + 5), fill="#e85d3f", width=16)
    draw.line((cx - 35, cy + 52, cx + 14, cy + 52), fill="#e85d3f", width=16)
    draw.text((90, 1530), "RESENHA CRÍTICA AUTORAL", font=font_badge, fill="#ffffff")
    draw.text((90, 1658), "DIO  •  PORTFÓLIO  •  2026", font=font_meta, fill="#8fa2c2")

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, optimize=True)


def iter_inline(text: str) -> Iterable[tuple[str, str, str | None]]:
    """Produz tokens (tipo, texto, URL opcional) para formatação inline."""
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            yield "text", text[cursor : match.start()], None
        if match.group(1) is not None:
            yield "bold", match.group(1), None
        elif match.group(2) is not None:
            yield "italic", match.group(2), None
        elif match.group(3) is not None:
            yield "link", match.group(3), match.group(4)
        elif match.group(5) is not None:
            yield "link", match.group(5), match.group(5)
        else:
            yield "footnote", f"[{match.group(6)}]", None
        cursor = match.end()
    if cursor < len(text):
        yield "text", text[cursor:], None


def add_docx_hyperlink(paragraph, text: str, url: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_docx_inline(paragraph, text: str) -> None:
    for kind, value, url in iter_inline(text):
        if kind == "link":
            add_docx_hyperlink(paragraph, value, url or value)
            continue
        run = paragraph.add_run(value)
        if kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True
        elif kind == "footnote":
            run.font.superscript = True
            run.font.size = _docx_pt(8)


def _docx_pt(value: float):
    from docx.shared import Pt

    return Pt(value)


def set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def set_repeat_table_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_props.append(repeat)


def add_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def create_docx(blocks: list[Block], output: Path, cover: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    props = document.core_properties
    props.title = "Scrum no FBI: uma análise crítica do Projeto Sentinel"
    props.subject = "Resenha crítica — Scrum: Analisando Cases de Sucesso"
    props.author = "0Barone"
    props.keywords = "Scrum, Agile, FBI, Sentinel, DIO, estudo de caso"
    props.comments = "Conteúdo autoral, produzido a partir de fontes citadas."

    # Capa sem cabeçalho, com margens mínimas.
    cover_section = document.sections[0]
    cover_section.page_width = Cm(21)
    cover_section.page_height = Cm(29.7)
    cover_section.top_margin = Cm(0.5)
    cover_section.bottom_margin = Cm(0.5)
    cover_section.left_margin = Cm(0.5)
    cover_section.right_margin = Cm(0.5)
    cover_section.header_distance = Cm(0)
    cover_section.footer_distance = Cm(0)

    cover_paragraph = document.add_paragraph()
    cover_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_paragraph.paragraph_format.space_after = Pt(0)
    cover_paragraph.paragraph_format.line_spacing = Pt(1)
    cover_paragraph.add_run().add_picture(str(cover), width=Cm(20), height=Cm(28.3))

    content = document.add_section(WD_SECTION.NEW_PAGE)
    content.page_width = Cm(21)
    content.page_height = Cm(29.7)
    content.top_margin = Cm(2.2)
    content.bottom_margin = Cm(2.0)
    content.left_margin = Cm(2.5)
    content.right_margin = Cm(2.2)
    content.header_distance = Cm(0.9)
    content.footer_distance = Cm(0.9)
    content.header.is_linked_to_previous = False
    content.footer.is_linked_to_previous = False

    # Reinicia a numeração no conteúdo.
    page_number_type = OxmlElement("w:pgNumType")
    page_number_type.set(qn("w:start"), "1")
    content._sectPr.append(page_number_type)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)

    for style_name, size, color in [
        ("Title", 25, NAVY),
        ("Heading 1", 19, NAVY),
        ("Heading 2", 15, BLUE),
        ("Heading 3", 12, TEAL),
    ]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)

    if "Referência" not in [s.name for s in document.styles]:
        reference_style = document.styles.add_style("Referência", WD_STYLE_TYPE.PARAGRAPH)
        reference_style.base_style = document.styles["Normal"]
        reference_style.font.size = Pt(9)
        reference_style.font.color.rgb = RGBColor.from_string("3F4D66")
        reference_style.paragraph_format.left_indent = Cm(0.7)
        reference_style.paragraph_format.first_line_indent = Cm(-0.7)
        reference_style.paragraph_format.space_after = Pt(7)

    header = content.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("SCRUM NO FBI  •  PROJETO SENTINEL")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(SLATE)

    footer = content.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("DIO  •  PORTFÓLIO  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(SLATE)
    add_page_number(footer)

    # Abre o conteúdo com título e metadados editoriais.
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Scrum no FBI: uma análise crítica do Projeto Sentinel")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.paragraph_format.space_after = Pt(16)
    subrun = subtitle.add_run("Resenha crítica • Desafio DIO • Outubro de 2026")
    subrun.font.size = Pt(10)
    subrun.font.bold = True
    subrun.font.color.rgb = RGBColor.from_string(ORANGE)

    # Barra editorial.
    bar = document.add_table(rows=1, cols=1)
    bar.autofit = False
    bar.columns[0].width = Cm(16)
    cell = bar.cell(0, 0)
    set_cell_shading(cell, ORANGE)
    cell.height = Cm(0.08)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    document.add_paragraph().paragraph_format.space_after = Pt(0)

    in_references = False
    for block in blocks:
        if block.kind == "heading":
            if block.text == "Referências":
                in_references = True
            style = "Heading 1" if block.level == 1 else "Heading 2" if block.level == 2 else "Heading 3"
            paragraph = document.add_paragraph(style=style)
            add_docx_inline(paragraph, block.text)
            continue

        if block.kind == "paragraph":
            paragraph = document.add_paragraph(style="Referência" if in_references else "Normal")
            add_docx_inline(paragraph, block.text)
            continue

        if block.kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(-0.35)
            paragraph.paragraph_format.space_after = Pt(4)
            add_docx_inline(paragraph, block.text)
            continue

        if block.kind == "number":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(-0.65)
            paragraph.paragraph_format.space_after = Pt(5)
            number = paragraph.add_run(f"{block.number}. ")
            number.bold = True
            number.font.color.rgb = RGBColor.from_string(BLUE)
            add_docx_inline(paragraph, block.text)
            continue

        if block.kind == "reference":
            paragraph = document.add_paragraph(style="Referência")
            lead = paragraph.add_run(f"[{block.number}] ")
            lead.bold = True
            lead.font.color.rgb = RGBColor.from_string(BLUE)
            add_docx_inline(paragraph, block.text)
            continue

        if block.kind == "rule":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(8)
            p_props = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "DCE2EE")
            borders.append(bottom)
            p_props.append(borders)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def rl_inline(text: str) -> str:
    chunks: list[str] = []
    for kind, value, url in iter_inline(text):
        safe = html.escape(value, quote=True)
        if kind == "bold":
            chunks.append(f"<b>{safe}</b>")
        elif kind == "italic":
            chunks.append(f"<i>{safe}</i>")
        elif kind == "link":
            safe_url = html.escape(url or value, quote=True)
            chunks.append(f'<a href="{safe_url}" color="#{BLUE}">{safe}</a>')
        elif kind == "footnote":
            chunks.append(f"<super>{safe}</super>")
        else:
            chunks.append(safe)
    return "".join(chunks)


def create_pdf(blocks: list[Block], output: Path) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            HRFlowable,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ImportError as exc:
        raise SystemExit(
            "ReportLab não está instalado. Execute: python -m pip install -r requirements.txt"
        ) from exc

    regular = find_font("DejaVuSans.ttf", "LiberationSans-Regular.ttf")
    bold = find_font("DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf")
    # Alguns ambientes mínimos oferecem apenas os pesos regular e bold.
    # Neles, o itálico usa a mesma fonte para manter a geração portátil.
    try:
        italic = find_font("DejaVuSans-Oblique.ttf", "LiberationSans-Italic.ttf")
    except FileNotFoundError:
        italic = regular
    try:
        bold_italic = find_font("DejaVuSans-BoldOblique.ttf", "LiberationSans-BoldItalic.ttf")
    except FileNotFoundError:
        bold_italic = bold
    pdfmetrics.registerFont(TTFont("ProjectSans", str(regular)))
    pdfmetrics.registerFont(TTFont("ProjectSans-Bold", str(bold)))
    pdfmetrics.registerFont(TTFont("ProjectSans-Italic", str(italic)))
    pdfmetrics.registerFont(TTFont("ProjectSans-BoldItalic", str(bold_italic)))
    pdfmetrics.registerFontFamily(
        "ProjectSans",
        normal="ProjectSans",
        bold="ProjectSans-Bold",
        italic="ProjectSans-Italic",
        boldItalic="ProjectSans-BoldItalic",
    )

    page_width, page_height = A4
    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=2.35 * cm,
        rightMargin=2.15 * cm,
        topMargin=2.25 * cm,
        bottomMargin=2.0 * cm,
        title="Scrum no FBI: uma análise crítica do Projeto Sentinel",
        author="0Barone",
        subject="Desafio DIO — Scrum: Analisando Cases de Sucesso",
        keywords="Scrum, Agile, FBI, Sentinel, DIO",
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ProjectBody",
        parent=styles["BodyText"],
        fontName="ProjectSans",
        fontSize=10.2,
        leading=14.8,
        textColor=colors.HexColor(f"#{NAVY}"),
        alignment=TA_JUSTIFY,
        spaceAfter=7,
        splitLongWords=True,
    )
    h1 = ParagraphStyle(
        "ProjectH1",
        parent=styles["Heading1"],
        fontName="ProjectSans-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor(f"#{NAVY}"),
        spaceBefore=15,
        spaceAfter=8,
        keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "ProjectH2",
        parent=styles["Heading2"],
        fontName="ProjectSans-Bold",
        fontSize=13.5,
        leading=17,
        textColor=colors.HexColor(f"#{BLUE}"),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True,
    )
    h3 = ParagraphStyle(
        "ProjectH3",
        parent=styles["Heading3"],
        fontName="ProjectSans-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor(f"#{TEAL}"),
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True,
    )
    bullet = ParagraphStyle(
        "ProjectBullet",
        parent=body,
        leftIndent=15,
        firstLineIndent=-11,
        spaceAfter=4,
        bulletIndent=0,
    )
    reference = ParagraphStyle(
        "ProjectReference",
        parent=body,
        fontSize=8.2,
        leading=11.5,
        leftIndent=16,
        firstLineIndent=-16,
        textColor=colors.HexColor("#3F4D66"),
        spaceAfter=7,
        alignment=TA_LEFT,
        splitLongWords=True,
        wordWrap="CJK",
    )
    title_style = ParagraphStyle(
        "ProjectTitle",
        parent=h1,
        fontSize=22,
        leading=27,
        spaceAfter=7,
    )
    meta_style = ParagraphStyle(
        "ProjectMeta",
        parent=body,
        fontName="ProjectSans-Bold",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor(f"#{ORANGE}"),
        alignment=TA_LEFT,
        spaceAfter=14,
    )

    def draw_cover(canvas, _doc):
        canvas.saveState()
        canvas.setTitle("Scrum no FBI: uma análise crítica do Projeto Sentinel")
        canvas.setAuthor("0Barone")
        canvas.setSubject("Resenha crítica autoral para o desafio da DIO")
        canvas.setFillColor(colors.HexColor("#10182E"))
        canvas.rect(0, 0, page_width, page_height, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#20335A"))
        canvas.circle(page_width + 1.0 * cm, page_height - 2.8 * cm, 5.8 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#101A34"))
        canvas.circle(page_width + 1.5 * cm, 2.4 * cm, 6.8 * cm, stroke=0, fill=1)

        canvas.setFillColor(colors.HexColor(f"#{TEAL}"))
        canvas.roundRect(2.0 * cm, page_height - 3.6 * cm, 5.3 * cm, 0.75 * cm, 0.35 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("ProjectSans-Bold", 10)
        canvas.drawCentredString(4.65 * cm, page_height - 3.35 * cm, "ESTUDO DE CASO")

        canvas.setFillColor(colors.HexColor("#FF8A5B"))
        canvas.setFont("ProjectSans-Bold", 11)
        canvas.drawString(2.0 * cm, page_height - 6.0 * cm, "SCRUM  •  AGILE  •  GOVERNO")
        canvas.setFillColor(colors.white)
        canvas.setFont("ProjectSans-Bold", 31)
        canvas.drawString(2.0 * cm, page_height - 8.2 * cm, "Scrum no FBI")
        canvas.setFillColor(colors.HexColor("#D6E0F3"))
        canvas.setFont("ProjectSans", 21)
        canvas.drawString(2.0 * cm, page_height - 9.7 * cm, "Projeto Sentinel")
        canvas.setFillColor(colors.HexColor(f"#{ORANGE}"))
        canvas.roundRect(2.0 * cm, page_height - 10.8 * cm, 12.2 * cm, 0.08 * cm, 0.04 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#B9C7DF"))
        canvas.setFont("ProjectSans", 12)
        canvas.drawString(2.0 * cm, page_height - 12.2 * cm, "Uma análise crítica sobre agilidade,")
        canvas.drawString(2.0 * cm, page_height - 12.85 * cm, "governança e transformação digital")

        # Ícone vetorial próprio.
        cx, cy = 15.9 * cm, 10.5 * cm
        canvas.setFillColor(colors.HexColor("#F7F8FB"))
        canvas.circle(cx, cy, 2.65 * cm, stroke=0, fill=1)
        canvas.setStrokeColor(colors.HexColor(f"#{NAVY}"))
        canvas.setLineWidth(5)
        canvas.roundRect(cx - 1.0 * cm, cy - 1.45 * cm, 2.0 * cm, 2.9 * cm, 0.85 * cm, stroke=1, fill=0)
        canvas.setStrokeColor(colors.HexColor(f"#{ORANGE}"))
        canvas.setLineWidth(4)
        for y_offset, x2 in [(0.75, 0.65), (0.0, 0.65), (-0.75, 0.28)]:
            canvas.line(cx - 0.52 * cm, cy + y_offset * cm, cx + x2 * cm, cy + y_offset * cm)

        canvas.setFillColor(colors.white)
        canvas.setFont("ProjectSans-Bold", 12)
        canvas.drawString(2.0 * cm, 2.8 * cm, "RESENHA CRÍTICA AUTORAL")
        canvas.setFillColor(colors.HexColor("#7F94B8"))
        canvas.setFont("ProjectSans", 9)
        canvas.drawString(2.0 * cm, 1.8 * cm, "DIO  •  PORTFÓLIO  •  2026")
        canvas.restoreState()

    def draw_content_page(canvas, _doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#DCE2EE"))
        canvas.setLineWidth(0.5)
        canvas.line(2.35 * cm, page_height - 1.45 * cm, page_width - 2.15 * cm, page_height - 1.45 * cm)
        canvas.setFillColor(colors.HexColor(f"#{SLATE}"))
        canvas.setFont("ProjectSans-Bold", 7.7)
        canvas.drawString(2.35 * cm, page_height - 1.18 * cm, "SCRUM NO FBI  •  PROJETO SENTINEL")
        logical_page = canvas.getPageNumber() - 1
        canvas.setFont("ProjectSans", 8)
        canvas.drawCentredString(page_width / 2, 1.1 * cm, f"DIO  •  PORTFÓLIO  •  {logical_page}")
        canvas.restoreState()

    story = [Spacer(1, 1), PageBreak()]
    story.append(Paragraph("Scrum no FBI: uma análise crítica do Projeto Sentinel", title_style))
    story.append(Paragraph("Resenha crítica • Desafio DIO • Outubro de 2026", meta_style))
    story.append(HRFlowable(width="100%", thickness=2.2, color=colors.HexColor(f"#{ORANGE}"), spaceAfter=12))

    in_references = False
    for block in blocks:
        if block.kind == "heading":
            if block.text == "Referências":
                in_references = True
            style = h1 if block.level == 1 else h2 if block.level == 2 else h3
            story.append(Paragraph(rl_inline(block.text), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(rl_inline(block.text), reference if in_references else body))
        elif block.kind == "bullet":
            story.append(Paragraph(f"•&nbsp;&nbsp;{rl_inline(block.text)}", bullet))
        elif block.kind == "number":
            story.append(Paragraph(f"<b>{block.number}.</b>&nbsp;&nbsp;{rl_inline(block.text)}", bullet))
        elif block.kind == "reference":
            story.append(Paragraph(f"<b>[{block.number}]</b>&nbsp;&nbsp;{rl_inline(block.text)}", reference))
        elif block.kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#DCE2EE"), spaceBefore=8, spaceAfter=8))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story, onFirstPage=draw_cover, onLaterPages=draw_content_page)


def main() -> None:
    blocks = parse_markdown(SOURCE)
    if not blocks or blocks[0].text != "Resumo":
        raise SystemExit("A estrutura esperada da resenha não foi encontrada.")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    generate_cover_png(COVER_PNG)
    create_docx(blocks, DOCX_PATH, COVER_PNG)
    create_pdf(blocks, PDF_PATH)
    print(f"Gerado: {DOCX_PATH.relative_to(ROOT)}")
    print(f"Gerado: {PDF_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
